#!/usr/bin/env python3
"""Check a deck built for Research Slides, without any agent-specific tooling.

  check-deck.py deck.pptx            # structural checks (needs python-pptx)
  check-deck.py deck.pptx --render   # also render every slide to PNG with LibreOffice, if installed
  check-deck.py deck.pptx --render --out preview/   # where the PNGs go (default: <deck>-preview/)

Structural checks (python-pptx): every shape inside the 10 x 5.625 in slide, no text below 8pt,
no empty placeholders, no leftover "Click to add" text, the footnote box narrower than 8.6 in,
each slide's title/topic present on content slides. Rendering (soffice --headless -> PDF ->
pdftoppm) is how an agent without a slide viewer LOOKS at the result: open the PNGs and read them
as the audience would. python-pptx: `pip install python-pptx` or `uv run --with python-pptx`.
"""
import sys, os, re, shutil, subprocess, argparse, glob

W_IN, H_IN = 10.0, 5.625
EMU = 914400

def structural(path):
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        print("python-pptx not installed: `pip install python-pptx` (or run with `uv run --with python-pptx`). Skipping structural checks.")
        return 0
    prs = Presentation(path)
    problems = 0
    def warn(i, msg):
        nonlocal problems; problems += 1; print(f"  slide {i}: {msg}")
    print(f"{os.path.basename(path)}: {len(prs.slides)} slides, {prs.slide_width/EMU:.2f} x {prs.slide_height/EMU:.2f} in")
    for i, slide in enumerate(prs.slides, 1):
        layout = slide.slide_layout.name
        texts = []
        for sh in slide.shapes:
            if sh.left is None or sh.width is None:
                continue
            l, t, w, h = (v / EMU for v in (sh.left, sh.top, sh.width, sh.height))
            is_num = 'slide number' in sh.name.lower() or (sh.is_placeholder and 'SLIDE_NUMBER' in str(sh.placeholder_format.type))
            if is_num:
                continue  # the number box overhangs the right edge by design (centred text, from the source decks)
            if l < -0.01 or t < -0.02 or l + w > W_IN + 0.01 or t + h > H_IN + 0.02:
                warn(i, f"'{sh.name}' leaves the slide ({l:.2f},{t:.2f}) {w:.2f}x{h:.2f} in")
            if sh.is_placeholder and sh.has_text_frame and not sh.text_frame.text.strip():
                warn(i, f"empty placeholder '{sh.name}'")
            if sh.has_text_frame:
                tx = sh.text_frame.text
                texts.append(tx)
                if 'Click to add' in tx or 'Lorem ipsum' in tx:
                    warn(i, f"template text left in '{sh.name}'")
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size is not None and r.font.size.pt < 8 and r.text.strip():
                            warn(i, f"{r.font.size.pt:g}pt text in '{sh.name}': {r.text[:40]!r}")
                if sh.name.lower().startswith('footnote') or (t > 5.2 and l < 0.05 and w > 8.6):
                    if w > 8.6 + 0.01:
                        warn(i, f"footnote box {w:.2f} in wide reaches the slide number")
            if sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        for p in cell.text_frame.paragraphs:
                            for r in p.runs:
                                if r.font.size is not None and r.font.size.pt < 8 and r.text.strip():
                                    warn(i, f"{r.font.size.pt:g}pt table text: {r.text[:30]!r}")
        if layout in ('CONTENT', 'CONTENT_LEAD'):
            # the title is the placeholder of type TITLE, or failing that any text in the header band (top < 0.7 in)
            has_title = any(sh.has_text_frame and sh.text_frame.text.strip() and (
                (sh.is_placeholder and 'TITLE' in str(sh.placeholder_format.type)) or (sh.top is not None and sh.top / EMU < 0.7))
                for sh in slide.shapes)
            if not has_title:
                warn(i, "content slide without a title")
    print("structural: OK" if problems == 0 else f"structural: {problems} problem(s)")
    return problems

VERSIONED = re.compile(r'^libreoffice(\d+(?:\.\d+)*)$')  # libreoffice25.8, libreoffice26.8, ...

def versioned_soffice():
    """Every versioned LibreOffice on PATH or under /opt, newest first: [(version tuple, path)]."""
    found = {}
    for d in os.environ.get('PATH', '').split(os.pathsep) + ['/opt']:
        try: names = os.listdir(d)
        except OSError: continue
        for n in names:
            m = VERSIONED.match(n)
            if not m: continue
            p = os.path.join(d, n, 'program', 'soffice') if d == '/opt' else os.path.join(d, n)
            if os.path.isfile(p) and os.access(p, os.X_OK):
                found.setdefault(tuple(int(x) for x in m.group(1).split('.')), p)
    return sorted(found.items(), reverse=True)

def find_soffice():
    """PATH first (soffice, libreoffice, then the newest libreoffice<version>), the official /opt build,
    distro paths, the macOS app, then a flatpak wrapper."""
    for name in ('soffice', 'libreoffice'):
        p = shutil.which(name)
        if p: return p
    v = versioned_soffice()
    if v: return v[0][1]
    cands = ['/usr/lib/libreoffice/program/soffice', '/usr/lib64/libreoffice/program/soffice',
        '/Applications/LibreOffice.app/Contents/MacOS/soffice', 'C:/Program Files/LibreOffice/program/soffice.exe']
    for c in cands:
        if os.path.isfile(c) and os.access(c, os.X_OK): return c
    if shutil.which('flatpak') and subprocess.run(['flatpak', 'info', 'org.libreoffice.LibreOffice'], capture_output=True).returncode == 0:
        return 'flatpak run org.libreoffice.LibreOffice'
    return None

def render(path, out):
    soffice = find_soffice()
    if not soffice:
        print("LibreOffice not found: install it (deb/rpm, the official /opt/libreoffice* build, the macOS app, or the flatpak) to render slides.")
        return 1
    os.makedirs(out, exist_ok=True)
    subprocess.run(soffice.split() + ['--headless', '--convert-to', 'pdf', '--outdir', out, path], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
    pdf = os.path.join(out, os.path.splitext(os.path.basename(path))[0] + '.pdf')
    if shutil.which('pdftoppm'):
        subprocess.run(['pdftoppm', '-png', '-r', '96', pdf, os.path.join(out, 'slide')], check=True, timeout=300)
        pngs = sorted(glob.glob(os.path.join(out, 'slide-*.png')))
        print(f"rendered {len(pngs)} slides -> {out}/slide-NN.png (read them; the PDF is {pdf})")
    else:
        print(f"rendered -> {pdf} (install poppler's pdftoppm for per-slide PNGs)")
    return 0

if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('deck'); ap.add_argument('--render', action='store_true'); ap.add_argument('--out')
    a = ap.parse_args()
    n = structural(a.deck)
    if a.render:
        render(a.deck, a.out or os.path.splitext(a.deck)[0] + '-preview')
    sys.exit(1 if n else 0)
